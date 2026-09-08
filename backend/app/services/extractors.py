"""Per-format text extractors. Every format maps to ExtractedPage units so
chunking, embedding, citations, and page limits stay format-agnostic."""
from collections.abc import Callable
from io import BytesIO
from typing import Protocol

from app.models import Document
from app.services.preview import PREVIEW_PDF_FORMATS
from app.services.processing import ExtractedPage, NoExtractableTextError, UnsupportedFileError

SECTION_TARGET_WORDS = 800


class PreviewConverter(Protocol):
    def convert_to_pdf(self, data: bytes, document_format: str) -> bytes:
        pass


def _word_sections(text: str) -> list[ExtractedPage]:
    words = text.split()
    if not words:
        raise NoExtractableTextError("no extractable text")
    return [
        ExtractedPage(page_number=index + 1, text=" ".join(words[start : start + SECTION_TARGET_WORDS]))
        for index, start in enumerate(range(0, len(words), SECTION_TARGET_WORDS))
    ]


def _require_text(pages: list[ExtractedPage]) -> list[ExtractedPage]:
    if not any(page.text.strip() for page in pages):
        # Carry the page count so the max-pages cap can still be enforced for
        # image-only documents.
        raise NoExtractableTextError("no extractable text", page_count=len(pages))
    return pages


def extract_pdf(data: bytes) -> list[ExtractedPage]:
    import fitz

    try:
        pdf = fitz.open(stream=data, filetype="pdf")
    except Exception as error:
        raise UnsupportedFileError("This PDF file could not be opened. It may be corrupt.") from error
    with pdf:
        pages = [
            ExtractedPage(page_number=index, text=page.get_text("text"))
            for index, page in enumerate(pdf, start=1)
        ]
    return _require_text(pages)


def _docx_table_text(table) -> str:
    return "\n".join(
        " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
        for row in table.rows
        if any(cell.text.strip() for cell in row.cells)
    )


def extract_docx(data: bytes) -> list[ExtractedPage]:
    from docx import Document as DocxDocument
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        docx_document = DocxDocument(BytesIO(data))
    except Exception as error:
        raise UnsupportedFileError("This DOCX file could not be opened. It may be corrupt.") from error

    # Walk body blocks in document order so tables keep their place between
    # paragraphs. Split on explicit/rendered page breaks when the document has
    # them, otherwise fall back to fixed-size word sections.
    sections: list[str] = []
    current: list[str] = []
    for child in docx_document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            paragraph = Paragraph(child, docx_document)
            if paragraph.text.strip():
                current.append(paragraph.text)
            if child.xpath(".//w:br[@w:type='page'] | .//w:lastRenderedPageBreak"):
                sections.append("\n".join(current))
                current = []
        elif child.tag == qn("w:tbl"):
            table_text = _docx_table_text(Table(child, docx_document))
            if table_text:
                current.append(table_text)
    sections.append("\n".join(current))
    sections = [section for section in sections if section.strip()]

    if len(sections) > 1:
        return [ExtractedPage(page_number=index, text=section) for index, section in enumerate(sections, start=1)]
    return _word_sections("\n".join(sections))


def extract_pptx(data: bytes) -> list[ExtractedPage]:
    from pptx import Presentation

    try:
        presentation = Presentation(BytesIO(data))
    except Exception as error:
        raise UnsupportedFileError("This PPTX file could not be opened. It may be corrupt.") from error

    def slide_text(slide) -> str:
        parts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text)
            elif getattr(shape, "has_table", False) and shape.has_table:
                rows = [
                    " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    for row in shape.table.rows
                    if any(cell.text.strip() for cell in row.cells)
                ]
                if rows:
                    parts.append("\n".join(rows))
        return "\n".join(parts)

    pages = [
        ExtractedPage(page_number=index, text=slide_text(slide))
        for index, slide in enumerate(presentation.slides, start=1)
    ]
    return _require_text(pages)


def extract_txt(data: bytes) -> list[ExtractedPage]:
    return _word_sections(data.decode("utf-8", errors="replace"))


def extract_rtf(data: bytes) -> list[ExtractedPage]:
    from striprtf.striprtf import rtf_to_text

    try:
        text = rtf_to_text(data.decode("utf-8", errors="replace"))
    except Exception as error:
        raise UnsupportedFileError("This RTF file could not be read. It may be corrupt.") from error
    return _word_sections(text)


EXTRACTORS: dict[str, Callable[[bytes], list[ExtractedPage]]] = {
    "pdf": extract_pdf,
    "docx": extract_docx,
    "pptx": extract_pptx,
    "txt": extract_txt,
    "rtf": extract_rtf,
}


class DocumentTextExtractor:
    """Storage-backed extractor that picks the format extractor per document."""

    def __init__(self, storage_service=None, preview_converter: PreviewConverter | None = None):
        self.storage_service = storage_service
        self.preview_converter = preview_converter

    def extract_pages(self, document: Document) -> list[ExtractedPage]:
        if self.storage_service is None:
            return []

        file_bytes = self.storage_service.download_pdf(document)
        if not file_bytes:
            return []

        document_format = document.format or "pdf"
        try:
            pages = self._extract_pages(document, file_bytes, document_format)
        except NoExtractableTextError as error:
            if error.page_count is not None:
                document.page_count = error.page_count
            raise
        document.page_count = len(pages)
        return pages

    def _extract_pages(self, document: Document, file_bytes: bytes, document_format: str) -> list[ExtractedPage]:
        if document_format in PREVIEW_PDF_FORMATS and self.preview_converter is not None:
            preview_pdf = self.preview_converter.convert_to_pdf(file_bytes, document_format)
            try:
                self.storage_service.upload_preview_pdf(document, preview_pdf)
            except Exception as error:
                raise UnsupportedFileError("Document preview PDF could not be stored.") from error
            return extract_pdf(preview_pdf)

        extract = EXTRACTORS.get(document_format)
        if extract is None:
            raise UnsupportedFileError(f"No extractor is registered for format '{document_format}'.")
        return extract(file_bytes)
