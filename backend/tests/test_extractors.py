from pathlib import Path

import pytest

from app.services.extractors import (
    EXTRACTORS,
    extract_docx,
    extract_pdf,
    extract_pptx,
    extract_rtf,
    extract_txt,
)
from app.services.processing import ExtractedPage, NoExtractableTextError, UnsupportedFileError

FIXTURES = Path(__file__).parent / "fixtures"


def _fixture_bytes(name: str) -> bytes:
    return (FIXTURES / name).read_bytes()


def test_registry_covers_all_supported_formats():
    assert set(EXTRACTORS) == {"pdf", "docx", "pptx", "txt", "rtf"}


def test_extract_docx_splits_on_explicit_page_breaks():
    pages = extract_docx(_fixture_bytes("sample.docx"))

    assert [page.page_number for page in pages] == [1, 2]
    assert "first page with useful text" in pages[0].text
    assert "second page after an explicit page break" in pages[1].text


def test_extract_docx_without_page_breaks_splits_by_word_count():
    from docx import Document as DocxDocument
    from io import BytesIO

    buffer = BytesIO()
    docx_document = DocxDocument()
    docx_document.add_paragraph(" ".join(f"word{index}" for index in range(1000)))
    docx_document.save(buffer)

    pages = extract_docx(buffer.getvalue())

    assert [page.page_number for page in pages] == [1, 2]
    assert len(pages[0].text.split()) == 800
    assert len(pages[1].text.split()) == 200


def test_extract_pptx_produces_one_page_per_slide():
    pages = extract_pptx(_fixture_bytes("sample.pptx"))

    assert [page.page_number for page in pages] == [1, 2]
    assert "slide one" in pages[0].text
    assert "slide two" in pages[1].text


def test_extract_txt_splits_every_800_words():
    words = " ".join(f"term{index}" for index in range(900))

    pages = extract_txt(words.encode("utf-8"))

    assert [page.page_number for page in pages] == [1, 2]
    assert pages[0].text.split()[0] == "term0"
    assert len(pages[0].text.split()) == 800


def test_extract_txt_fixture_yields_single_section():
    pages = extract_txt(_fixture_bytes("sample.txt"))

    assert pages == [ExtractedPage(page_number=1, text=pages[0].text)]
    assert "Fixture TXT file" in pages[0].text


def test_extract_rtf_strips_markup_to_text():
    pages = extract_rtf(_fixture_bytes("sample.rtf"))

    assert len(pages) == 1
    assert "Fixture RTF file with useful text." in pages[0].text
    assert "\\rtf1" not in pages[0].text


@pytest.mark.parametrize("format_key", ["docx", "pptx", "txt", "rtf"])
def test_empty_file_raises_no_extractable_text(format_key):
    extract = EXTRACTORS[format_key]

    if format_key in ("docx", "pptx"):
        # An empty byte string is not even a valid package for OOXML formats.
        with pytest.raises(UnsupportedFileError):
            extract(b"")
    else:
        with pytest.raises(NoExtractableTextError):
            extract(b"")


def test_whitespace_only_txt_raises_no_extractable_text():
    with pytest.raises(NoExtractableTextError):
        extract_txt(b"   \n \t ")


def test_garbage_bytes_as_docx_raises_unsupported_file_error():
    with pytest.raises(UnsupportedFileError):
        extract_docx(b"this is definitely not a zip archive")


def test_extract_pdf_no_text_error_carries_page_count():
    import fitz

    pdf = fitz.open()
    for _ in range(3):
        pdf.new_page()

    with pytest.raises(NoExtractableTextError) as exc_info:
        extract_pdf(pdf.tobytes())

    assert exc_info.value.page_count == 3


def test_extract_docx_includes_table_text():
    from docx import Document as DocxDocument
    from io import BytesIO

    doc = DocxDocument()
    doc.add_paragraph("Rates overview.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Plan"
    table.cell(0, 1).text = "Rate"
    table.cell(1, 0).text = "Standard"
    table.cell(1, 1).text = "42 dollars"
    doc.add_paragraph("End of document.")
    buffer = BytesIO()
    doc.save(buffer)

    pages = extract_docx(buffer.getvalue())

    text = "\n".join(page.text for page in pages)
    assert "Rates overview." in text
    assert "Standard | 42 dollars" in text
    assert text.index("Rates overview.") < text.index("Standard | 42 dollars") < text.index("End of document.")


def test_extract_pptx_includes_table_text():
    from io import BytesIO

    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(2))
    shape.table.cell(0, 0).text = "Quarter"
    shape.table.cell(0, 1).text = "Revenue"
    shape.table.cell(1, 0).text = "Q4"
    shape.table.cell(1, 1).text = "9 million"
    buffer = BytesIO()
    presentation.save(buffer)

    pages = extract_pptx(buffer.getvalue())

    assert len(pages) == 1
    assert "Q4 | 9 million" in pages[0].text

